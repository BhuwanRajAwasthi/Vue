from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path(__file__).parent / 'NepKart_Project_Documentation.docx'
doc = Document()

# Page setup
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)
for name, size, color in [('Title', 28, '17365D'), ('Heading 1', 18, '17365D'), ('Heading 2', 13, '2E6041'), ('Heading 3', 11, '404040')]:
    styles[name].font.name = 'Aptos Display'
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

if 'Code' not in styles:
    code_style = styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = 'Consolas'
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = RGBColor(55, 55, 55)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, 'FFFFFF')
        shade(t.rows[0].cells[i], '17365D')
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if len(t.rows) % 2 == 0:
                shade(cells[i], 'F3F6F9')
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return t


def bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.add_run(text)
    return p


def number(text):
    p = doc.add_paragraph(style='List Number')
    p.add_run(text)
    return p


def code(text):
    p = doc.add_paragraph(style='Code')
    p.paragraph_format.left_indent = Inches(0.25)
    p.add_run(text)
    return p


def page_break():
    doc.add_page_break()

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('\n\nNEPKART\n').bold = True
p.runs[0].font.size = Pt(34)
p.runs[0].font.color.rgb = RGBColor.from_string('17365D')
p.add_run('Full-Stack E-Commerce Web Application').bold = True
p.runs[-1].font.size = Pt(18)
p.runs[-1].font.color.rgb = RGBColor.from_string('2E6041')

doc.add_paragraph('\n')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Project Documentation\n').font.size = Pt(20)
p.add_run('Prepared for academic project presentation\n\n')
p.add_run('Technology: Vue 3, TypeScript, Node.js, Express, MongoDB').italic = True

doc.add_paragraph('\n')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Document scope: architecture, workflow, services, dependencies, database, security, setup, and operations')
page_break()

# Contents-like overview
doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph('NepKart is a full-stack online shopping platform. It allows customers to browse products, filter products by category and audience, view multiple product images, register and verify accounts, manage a cart, place orders, and choose Cash on Delivery or eSewa payment. Administrators manage the product catalogue, upload images, and update order states from an admin dashboard.')
doc.add_paragraph('The project uses a separate Vue frontend and Express backend. The frontend communicates with the backend through REST APIs. MongoDB stores users, products, orders, and cart data. Cloudinary stores uploaded image files while MongoDB stores their secure URLs.')

table(['Area', 'Implementation'], [
    ('Frontend', 'Vue 3 Composition API with TypeScript and Vite'),
    ('State management', 'Pinia stores for authentication, products, cart, and checkout'),
    ('Backend', 'Node.js and Express REST API'),
    ('Database', 'MongoDB accessed through Mongoose'),
    ('Authentication', 'JWT, bcrypt password hashing, email OTP, phone OTP, Google OAuth, GitHub OAuth'),
    ('Assets', 'Cloudinary for product and profile images'),
    ('Payments', 'Cash on Delivery and eSewa ePay v2 integration'),
    ('Notifications', 'Nodemailer SMTP for verification and order emails'),
])

# Architecture
doc.add_heading('2. System Architecture', level=1)
doc.add_paragraph('NepKart follows a three-tier architecture. The presentation layer runs in the browser, the application layer exposes REST endpoints and business rules, and the data layer stores persistent records in MongoDB. External services support image hosting, email, OAuth, SMS, and payment processing.')
code('''+------------------------- CLIENT BROWSER --------------------------+
| Vue 3 + TypeScript + Vite + Pinia + Vue Router + PrimeVue         |
| Storefront | Authentication | Cart | Checkout | Admin dashboard  |
+-----------------------------+------------------------------------+
                              | HTTP JSON / multipart requests
                              v
+------------------------- EXPRESS API ----------------------------+
| CORS | JSON parser | JWT middleware | Admin role middleware      |
| Auth | Products | Cart | Orders | Payments                      |
+-------------+----------------+----------------+---------------------+
              |                |                |
              v                v                v
        MongoDB/Mongoose   Cloudinary       Email/OAuth/eSewa
        users/products/    image URLs       external services
        orders/cart''')

doc.add_heading('2.1 Main Components', level=2)
table(['Component', 'Responsibility'], [
    ('Vue application', 'Renders pages, collects user input, calls APIs, and manages client state.'),
    ('Pinia authStore', 'Stores JWT and current user, performs login, registration, verification, profile updates, and logout.'),
    ('Pinia productStore', 'Loads products, filters and searches catalogue data, maintains cart state, saves cart, and starts checkout.'),
    ('Express routes', 'Validate requests and implement authentication, catalogue, cart, order, and payment operations.'),
    ('JWT middleware', 'Verifies bearer tokens and attaches the authenticated user to the request.'),
    ('Admin middleware', 'Restricts product and order-management operations to users with role=admin.'),
    ('Mongoose models', 'Define data structure, validation, relations, defaults, and password hashing hooks.'),
])

# Structure
doc.add_heading('3. Project Structure', level=1)
code('''Vue/
|-- Backend/
|   |-- index.js                 Express server entry point
|   |-- package.json             Backend dependencies and scripts
|   |-- config/cloudinary.js     Cloudinary configuration
|   |-- middleware/auth.js       JWT and admin authorization
|   |-- models/                  User, Product, and Order schemas
|   |-- routes/                  Auth, products, cart, orders, payments
|   |-- utils/                   Mailer and SMS helpers
|   |-- seed.js                  Admin and product seed script
|   `-- .env                     Local secrets and service configuration
|-- Frontend/
|   |-- package.json             Frontend dependencies and scripts
|   |-- src/App.vue              Application shell
|   |-- src/main.ts              Vue, Pinia, Router, PrimeVue bootstrap
|   |-- src/router.ts            Client routes and auth redirect
|   |-- src/stores/pinia/        Authentication and product/cart state
|   |-- src/views/               Storefront, auth, admin, cart, profile, orders
|   `-- public/                  Static frontend assets
`-- context/                      Project notes and this documentation''')

# Workflows
doc.add_heading('4. Core Workflows', level=1)
doc.add_heading('4.1 Customer Registration and Login', level=2)
for item in [
    'The user submits name, email, password, phone, address, and city from Register.vue.',
    'The backend normalizes the email, hashes the password through the User pre-save hook, creates a six-digit email verification code hash, and sends the code through SMTP.',
    'The user enters the code. The backend verifies the hash and expiry, marks emailVerified=true, and returns a JWT.',
    'For login, the backend checks the normalized email, verified status, and bcrypt password match, then returns JWT plus user profile data.',
    'The auth store saves the JWT in localStorage and stores the user in Pinia. The router redirects authenticated users away from login and registration pages.'
]: number(item)

doc.add_heading('4.2 Product and Image Workflow', level=2)
for item in [
    'An administrator selects one or more product images in the Admin Panel.',
    'The frontend sends multipart/form-data to POST /api/products/upload.',
    'Multer receives the files and the backend uploads them to Cloudinary under ecommerce-products.',
    'Cloudinary returns secure URLs. The frontend appends all URLs to the product image list and uses the first image as the main image.',
    'The product is saved in MongoDB with gender, category, image, images, price, stock, sizes, and colors.',
    'The home page and product detail page display the main image and thumbnails for additional images.'
]: number(item)

doc.add_heading('4.3 Cart and Checkout Workflow', level=2)
for item in [
    'A customer adds a product and optional size/color selections to the Pinia cart.',
    'The cart is persisted through PUT /api/cart for the authenticated user.',
    'Prepare checkout opens the checkout form and pre-fills saved profile data.',
    'The customer selects Cash on Delivery or eSewa and confirms the delivery details.',
    'The backend validates product stock and creates an order. Cash on Delivery completes immediately.',
    'For eSewa, the backend creates a pending order, generates the required HMAC-SHA256 signature, and returns payment fields. The frontend submits those fields to the eSewa gateway.',
    'After payment, eSewa redirects to the backend callback. The callback verifies the response signature, updates the order to paid/processing, and redirects the user to the frontend.'
]: number(item)

doc.add_heading('4.4 Admin Order Workflow', level=2)
for item in [
    'An admin logs in and is identified by role=admin.',
    'The Admin Panel loads all orders through GET /api/orders.',
    'Orders can be filtered by pending, processing, shipped, delivered, or cancelled.',
    'An admin changes a status through PUT /api/orders/:id/status.',
    'Customers see their own orders and status history in the Orders view.'
]: number(item)

# Services
doc.add_heading('5. Integrated Services', level=1)
table(['Service', 'Purpose', 'Configuration variables'], [
    ('MongoDB', 'Persistent users, products, orders, and carts.', 'MONGODB_URI'),
    ('Cloudinary', 'Stores product and profile images and returns HTTPS URLs.', 'CLOUDINARY_CLOUD_NAME, CLOUDINARY_API_KEY, CLOUDINARY_API_SECRET'),
    ('SMTP/Nodemailer', 'Sends verification and order confirmation emails.', 'SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS, SMTP_FROM'),
    ('Google OAuth', 'Optional social login and account creation.', 'GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET, GOOGLE_REDIRECT_URI'),
    ('GitHub OAuth', 'Optional social login and account creation.', 'GITHUB_CLIENT_ID, GITHUB_CLIENT_SECRET, GITHUB_CALLBACK_URL'),
    ('SMS provider', 'Phone OTP request and verification helper.', 'Provider variables in Backend/.env'),
    ('eSewa ePay v2', 'Online payment redirect, callback, and payment status update.', 'ESEWA_PRODUCT_CODE, ESEWA_SECRET, ESEWA_URL, success/failure URLs'),
])

doc.add_paragraph('Important operational note: eSewa callback URLs must be publicly reachable in a deployed environment. localhost callback URLs work only when the gateway can access the local machine, which is generally not possible from an external payment gateway. Test and production merchant credentials and gateway URLs must match each other.')

# API
doc.add_heading('6. API Summary', level=1)
table(['Method', 'Endpoint', 'Purpose', 'Access'], [
    ('POST', '/api/auth/register', 'Create an account and send email verification code.', 'Public'),
    ('POST', '/api/auth/login', 'Authenticate with email and password.', 'Public'),
    ('POST', '/api/auth/verify-email-otp', 'Verify email and issue JWT.', 'Public'),
    ('POST', '/api/auth/resend-email-verification', 'Send a fresh verification code.', 'Public'),
    ('POST', '/api/auth/phone/request-otp', 'Request phone login code.', 'Public'),
    ('POST', '/api/auth/phone/verify-otp', 'Verify phone code and issue JWT.', 'Public'),
    ('GET/PUT', '/api/auth/me', 'Read or update authenticated profile.', 'JWT'),
    ('POST', '/api/auth/me/avatar', 'Upload profile image to Cloudinary.', 'JWT'),
    ('GET', '/api/products', 'List products.', 'Public'),
    ('GET', '/api/products/:id', 'Get one product.', 'Public'),
    ('POST', '/api/products/upload', 'Upload up to eight product images.', 'JWT + Admin'),
    ('POST/PUT/DELETE', '/api/products[/:id]', 'Create, edit, or delete products.', 'JWT + Admin'),
    ('GET/PUT', '/api/cart', 'Read or replace authenticated cart.', 'JWT'),
    ('POST', '/api/orders', 'Create Cash on Delivery order.', 'JWT'),
    ('GET', '/api/orders', 'List own orders or all orders for admin.', 'JWT'),
    ('PUT', '/api/orders/:id/status', 'Update order state.', 'JWT + Admin'),
    ('POST', '/api/payments/esewa/initiate', 'Create pending order and payment fields.', 'JWT'),
    ('GET', '/api/payments/esewa/success', 'Verify successful eSewa response.', 'Gateway callback'),
    ('GET', '/api/payments/esewa/failure', 'Handle failed/cancelled payment.', 'Gateway callback'),
])

# Data model
doc.add_heading('7. Database Design', level=1)
doc.add_heading('7.1 User Collection', level=2)
doc.add_paragraph('Stores identity, authentication, contact information, role, profile image, preferred payment method, OTP data, and embedded cart entries. Passwords are never stored in plain text; bcrypt hashes are stored instead.')
doc.add_heading('7.2 Product Collection', level=2)
doc.add_paragraph('Stores title, description, category, gender audience (men, women, or unisex), price, stock, image URLs, sizes, colors, brand, and timestamps.')
doc.add_heading('7.3 Order Collection', level=2)
doc.add_paragraph('Stores the customer reference, ordered products and prices, quantity, shipping address, payment details, subtotal, delivery fee, total, status, and timestamps. Status values are pending, processing, shipped, delivered, and cancelled.')
doc.add_heading('7.4 Relationships', level=2)
bullet('User.cart.product references Product._id.')
bullet('Order.user references User._id.')
bullet('Order.items.product references Product._id.')
bullet('Mongoose populate() is used when returning cart, order, and product details.')

# Dependencies
doc.add_heading('8. Dependencies', level=1)
table(['Package', 'Layer', 'Why it is used'], [
    ('vue', 'Frontend', 'Reactive UI framework.'),
    ('typescript', 'Frontend', 'Static typing and safer development.'),
    ('vite', 'Frontend', 'Development server and production bundler.'),
    ('vue-router', 'Frontend', 'Client-side navigation.'),
    ('pinia', 'Frontend', 'Centralized reactive state management.'),
    ('primevue / primeicons', 'Frontend', 'UI components and icons.'),
    ('tailwindcss / postcss', 'Frontend', 'Utility styling and CSS processing.'),
    ('express', 'Backend', 'HTTP server and REST routing.'),
    ('mongoose', 'Backend', 'MongoDB ODM, schemas, validation, and populate.'),
    ('jsonwebtoken', 'Backend', 'JWT creation and verification.'),
    ('bcryptjs', 'Backend', 'Password hashing and comparison.'),
    ('multer', 'Backend', 'Multipart file upload handling.'),
    ('cloudinary', 'Backend', 'Remote image storage and delivery.'),
    ('nodemailer', 'Backend', 'SMTP email delivery.'),
    ('dotenv', 'Backend', 'Loads environment variables.'),
    ('cors', 'Backend', 'Allows frontend-backend cross-origin requests.'),
])

# Security
doc.add_heading('9. Security and Access Control', level=1)
bullet('JWT bearer tokens protect authenticated API requests.')
bullet('The requireAuth middleware validates the token and loads the current user.')
bullet('The requireAdmin middleware blocks catalogue management and order status changes for normal users.')
bullet('Passwords are hashed with bcrypt before persistence.')
bullet('Email OTP codes are stored as SHA-256 hashes and expire after a limited period.')
bullet('Payment responses are checked with HMAC-SHA256 before an order is marked paid.')
bullet('Secrets belong in Backend/.env and must not be committed or exposed in documentation.')
bullet('Production deployment should use HTTPS, restricted CORS, rotated credentials, rate limiting, and server-side validation for all payment and profile data.')

# Setup
doc.add_heading('10. Setup and Running the Project', level=1)
doc.add_heading('10.1 Prerequisites', level=2)
bullet('Node.js 20.19+ or Node.js 22.12+.')
bullet('MongoDB running locally or a reachable MongoDB connection string.')
bullet('Configured Cloudinary and SMTP credentials for image upload and email features.')
bullet('eSewa merchant/test credentials and publicly reachable callback URLs for payment testing.')
doc.add_heading('10.2 Install and Seed', level=2)
code('''cd Backend
npm install
node seed.js

cd ..\\Frontend
npm install''')
doc.add_heading('10.3 Run Development Servers', level=2)
code('''# Terminal 1
cd Backend
node index.js

# Terminal 2
cd Frontend
npm run dev''')
doc.add_heading('10.4 Build and Verify', level=2)
code('''cd Frontend
npm run build
# Runs vue-tsc type-check and Vite production build''')
doc.add_paragraph('The backend listens on port 4000 by default. Vite normally serves the frontend on port 5173. The frontend uses VITE_API_URL when provided, otherwise it falls back to http://localhost:4000/api.')

# Testing and limitations
doc.add_heading('11. Testing and Current Limitations', level=1)
bullet('The frontend production build and vue-tsc type-check have been run successfully during development.')
bullet('Backend route syntax checks pass with node --check.')
bullet('There is no automated unit-test suite currently defined in package.json.')
bullet('The eSewa integration depends on external gateway availability, merchant credentials, and reachable callback URLs. A gateway-side 404 cannot be repaired by frontend code alone.')
bullet('Existing products created before the gender field was added are treated as unisex until edited.')
bullet('The application uses localStorage for the JWT; a production deployment should consider secure HTTP-only cookies depending on the security model.')

# Conclusion
doc.add_heading('12. Conclusion', level=1)
doc.add_paragraph('NepKart demonstrates a complete full-stack commerce workflow: product discovery, audience and category filtering, multi-image catalogue management, account verification, role-based administration, cart persistence, order lifecycle tracking, email communication, cloud image storage, and payment gateway integration. The modular separation between Vue views/stores, Express routes/middleware, Mongoose models, and external services makes the system understandable and extendable for future features such as reviews, inventory reservations, delivery tracking, reporting, and automated testing.')

# Footer
for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('NepKart Project Documentation | Confidential project material').font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(110, 110, 110)

# Core properties
doc.core_properties.title = 'NepKart Project Documentation'
doc.core_properties.subject = 'Full-stack e-commerce architecture and implementation'
doc.core_properties.author = 'NepKart Project Team'
doc.core_properties.keywords = 'Vue, Express, MongoDB, e-commerce, architecture'
doc.save(OUT)
print(OUT)
